from __future__ import annotations

import csv
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock


RELEASE_DIRECTORY = Path(__file__).resolve().parents[1]
SOURCE_DIRECTORY = RELEASE_DIRECTORY / "src"
sys.path.insert(0, str(SOURCE_DIRECTORY))

from sbml_structural import (
    directional_reaction_records,
    exchange_records,
    load_model,
    transport_records,
)
from step1_target_selection import read_targets, validate_targets
from step4_intracellular_production import (
    COFACTOR_BOOTSTRAP,
    executable_environment,
    net_production_feasibility,
    run_menetools,
)
from step6_cross_species_donor import read_model_manifest
from workbook_writer import write_workbook

sys.path.insert(0, str(RELEASE_DIRECTORY))
import find_crossfeeding
import validate_expected_cases


SYNTHETIC_SBML = """<?xml version="1.0" encoding="UTF-8"?>
<sbml xmlns="http://www.sbml.org/sbml/level3/version1/core"
      xmlns:fbc="http://www.sbml.org/sbml/level3/version1/fbc/version2"
      level="3" version="1" fbc:required="false">
  <model id="M_test" name="Test bacterium strain 1" fbc:strict="true">
    <listOfCompartments>
      <compartment id="c" constant="true"/>
      <compartment id="e" constant="true"/>
    </listOfCompartments>
    <listOfSpecies>
      <species id="M_target__91__c__93__" name="Target" compartment="c"/>
      <species id="M_target__91__e__93__" name="Target" compartment="e"/>
      <species id="M_prec__91__c__93__" name="Precursor" compartment="c"/>
      <species id="M_prec__91__e__93__" name="Precursor" compartment="e"/>
    </listOfSpecies>
    <listOfParameters>
      <parameter id="LB_REV" value="-1000" constant="true"/>
      <parameter id="LB_ZERO" value="0" constant="true"/>
      <parameter id="UB_ZERO" value="0" constant="true"/>
      <parameter id="UB_POS" value="1000" constant="true"/>
    </listOfParameters>
    <listOfReactions>
      <reaction id="EX_target" reversible="true" fbc:lowerFluxBound="LB_REV" fbc:upperFluxBound="UB_POS">
        <listOfReactants><speciesReference species="M_target__91__e__93__" stoichiometry="1"/></listOfReactants>
      </reaction>
      <reaction id="T_target" reversible="false" fbc:lowerFluxBound="LB_ZERO" fbc:upperFluxBound="UB_POS">
        <listOfReactants><speciesReference species="M_target__91__c__93__" stoichiometry="1"/></listOfReactants>
        <listOfProducts><speciesReference species="M_target__91__e__93__" stoichiometry="1"/></listOfProducts>
      </reaction>
      <reaction id="EX_prec" reversible="false" fbc:lowerFluxBound="LB_REV" fbc:upperFluxBound="UB_ZERO">
        <listOfReactants><speciesReference species="M_prec__91__e__93__" stoichiometry="1"/></listOfReactants>
      </reaction>
      <reaction id="T_prec" reversible="false" fbc:lowerFluxBound="LB_ZERO" fbc:upperFluxBound="UB_POS">
        <listOfReactants><speciesReference species="M_prec__91__e__93__" stoichiometry="1"/></listOfReactants>
        <listOfProducts><speciesReference species="M_prec__91__c__93__" stoichiometry="1"/></listOfProducts>
      </reaction>
      <reaction id="PROD" reversible="false" fbc:lowerFluxBound="LB_ZERO" fbc:upperFluxBound="UB_POS">
        <listOfReactants><speciesReference species="M_prec__91__c__93__" stoichiometry="1"/></listOfReactants>
        <listOfProducts><speciesReference species="M_target__91__c__93__" stoichiometry="1"/></listOfProducts>
      </reaction>
    </listOfReactions>
  </model>
</sbml>
"""


SYNTHETIC_PERIPLASM_SBML = """<?xml version="1.0" encoding="UTF-8"?>
<sbml xmlns="http://www.sbml.org/sbml/level3/version1/core"
      xmlns:fbc="http://www.sbml.org/sbml/level3/version1/fbc/version2"
      level="3" version="1" fbc:required="false">
  <model id="M_periplasm" name="Periplasm test" fbc:strict="true">
    <listOfCompartments>
      <compartment id="c" constant="true"/>
      <compartment id="p" constant="true"/>
      <compartment id="e" constant="true"/>
    </listOfCompartments>
    <listOfSpecies>
      <species id="M_target__91__c__93__" compartment="c"/>
      <species id="M_target__91__p__93__" compartment="p"/>
      <species id="M_target__91__e__93__" compartment="e"/>
    </listOfSpecies>
    <listOfParameters>
      <parameter id="LB_REV" value="-1000" constant="true"/>
      <parameter id="UB_POS" value="1000" constant="true"/>
    </listOfParameters>
    <listOfReactions>
      <reaction id="T_cp" reversible="true" fbc:lowerFluxBound="LB_REV" fbc:upperFluxBound="UB_POS">
        <listOfReactants><speciesReference species="M_target__91__c__93__"/></listOfReactants>
        <listOfProducts><speciesReference species="M_target__91__p__93__"/></listOfProducts>
      </reaction>
      <reaction id="T_pe" reversible="true" fbc:lowerFluxBound="LB_REV" fbc:upperFluxBound="UB_POS">
        <listOfReactants><speciesReference species="M_target__91__p__93__"/></listOfReactants>
        <listOfProducts><speciesReference species="M_target__91__e__93__"/></listOfProducts>
      </reaction>
    </listOfReactions>
  </model>
</sbml>
"""


SYNTHETIC_REDOX_SBML = """<?xml version="1.0" encoding="UTF-8"?>
<sbml xmlns="http://www.sbml.org/sbml/level3/version1/core"
      xmlns:fbc="http://www.sbml.org/sbml/level3/version1/fbc/version2"
      level="3" version="1" fbc:required="false">
  <model id="M_redox" fbc:strict="true">
    <listOfCompartments>
      <compartment id="c" constant="true"/>
      <compartment id="e" constant="true"/>
    </listOfCompartments>
    <listOfSpecies>
      <species id="M_target__91__c__93__" compartment="c"/>
      <species id="M_prec__91__c__93__" compartment="c"/>
      <species id="M_prec__91__e__93__" compartment="e"/>
      <species id="M_nad__91__c__93__" compartment="c"/>
      <species id="M_nadh__91__c__93__" compartment="c"/>
      <species id="M_pyr__91__c__93__" compartment="c"/>
      <species id="M_so3__91__c__93__" compartment="c"/>
    </listOfSpecies>
    <listOfParameters>
      <parameter id="LB_ZERO" value="0" constant="true"/>
      <parameter id="UB_POS" value="1000" constant="true"/>
    </listOfParameters>
    <listOfReactions>
      <reaction id="T_prec" reversible="false" fbc:lowerFluxBound="LB_ZERO" fbc:upperFluxBound="UB_POS">
        <listOfReactants><speciesReference species="M_prec__91__e__93__"/></listOfReactants>
        <listOfProducts><speciesReference species="M_prec__91__c__93__"/></listOfProducts>
      </reaction>
      <reaction id="LDH" reversible="false" fbc:lowerFluxBound="LB_ZERO" fbc:upperFluxBound="UB_POS">
        <listOfReactants>
          <speciesReference species="M_prec__91__c__93__"/>
          <speciesReference species="M_nad__91__c__93__"/>
        </listOfReactants>
        <listOfProducts>
          <speciesReference species="M_pyr__91__c__93__"/>
          <speciesReference species="M_nadh__91__c__93__"/>
        </listOfProducts>
      </reaction>
      <reaction id="REDUCTASE" reversible="false" fbc:lowerFluxBound="LB_ZERO" fbc:upperFluxBound="UB_POS">
        <listOfReactants>
          <speciesReference species="M_nadh__91__c__93__"/>
          <speciesReference species="M_so3__91__c__93__"/>
        </listOfReactants>
        <listOfProducts>
          <speciesReference species="M_nad__91__c__93__"/>
          <speciesReference species="M_target__91__c__93__"/>
        </listOfProducts>
      </reaction>
    </listOfReactions>
  </model>
</sbml>
"""


SYNTHETIC_TARGET_CYCLE_SBML = """<?xml version="1.0" encoding="UTF-8"?>
<sbml xmlns="http://www.sbml.org/sbml/level3/version1/core"
      xmlns:fbc="http://www.sbml.org/sbml/level3/version1/fbc/version2"
      level="3" version="1" fbc:required="false">
  <model id="M_target_cycle" fbc:strict="true">
    <listOfCompartments>
      <compartment id="c" constant="true"/>
      <compartment id="e" constant="true"/>
    </listOfCompartments>
    <listOfSpecies>
      <species id="M_target__91__c__93__" compartment="c"/>
      <species id="M_prec__91__c__93__" compartment="c"/>
      <species id="M_prec__91__e__93__" compartment="e"/>
      <species id="M_seed__91__c__93__" compartment="c"/>
      <species id="M_intermediate__91__c__93__" compartment="c"/>
    </listOfSpecies>
    <listOfParameters>
      <parameter id="LB_ZERO" value="0" constant="true"/>
      <parameter id="UB_POS" value="1000" constant="true"/>
    </listOfParameters>
    <listOfReactions>
      <reaction id="T_prec" reversible="false" fbc:lowerFluxBound="LB_ZERO" fbc:upperFluxBound="UB_POS">
        <listOfReactants><speciesReference species="M_prec__91__e__93__"/></listOfReactants>
        <listOfProducts><speciesReference species="M_prec__91__c__93__"/></listOfProducts>
      </reaction>
      <reaction id="BASE" reversible="false" fbc:lowerFluxBound="LB_ZERO" fbc:upperFluxBound="UB_POS">
        <listOfReactants><speciesReference species="M_seed__91__c__93__"/></listOfReactants>
        <listOfProducts><speciesReference species="M_target__91__c__93__"/></listOfProducts>
      </reaction>
      <reaction id="TARGET_CONSUMING_STEP" reversible="false" fbc:lowerFluxBound="LB_ZERO" fbc:upperFluxBound="UB_POS">
        <listOfReactants>
          <speciesReference species="M_prec__91__c__93__"/>
          <speciesReference species="M_target__91__c__93__"/>
        </listOfReactants>
        <listOfProducts><speciesReference species="M_intermediate__91__c__93__"/></listOfProducts>
      </reaction>
      <reaction id="TARGET_REFORMING_STEP" reversible="false" fbc:lowerFluxBound="LB_ZERO" fbc:upperFluxBound="UB_POS">
        <listOfReactants><speciesReference species="M_intermediate__91__c__93__"/></listOfReactants>
        <listOfProducts><speciesReference species="M_target__91__c__93__"/></listOfProducts>
      </reaction>
    </listOfReactions>
  </model>
</sbml>
"""


TARGET_HEADER = (
    "target_name\tbase_id\thost_relevance\tmetabolite_class\t"
    "microbial_evidence\tselection_rationale\treference\n"
)


class StructuralRulesTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.directory = Path(self.temporary.name)
        self.model_path = self.directory / "test.xml"
        self.model_path.write_text(SYNTHETIC_SBML, encoding="utf-8")

    def tearDown(self) -> None:
        self.temporary.cleanup()

    def test_exchange_and_transport_directions(self) -> None:
        model = load_model(self.model_path)
        target_exchange = exchange_records(model, "target")
        self.assertEqual(len(target_exchange), 1)
        self.assertTrue(target_exchange[0]["export_allowed"])
        self.assertTrue(target_exchange[0]["import_allowed"])
        target_transport = transport_records(model, "target")
        self.assertTrue(any(row["c_to_e_allowed"] for row in target_transport))
        self.assertFalse(any(row["e_to_c_allowed"] for row in target_transport))
        precursor_exchange = exchange_records(model, "prec")
        self.assertTrue(precursor_exchange[0]["import_allowed"])
        self.assertFalse(precursor_exchange[0]["export_allowed"])
        precursor_transport = transport_records(model, "prec")
        self.assertTrue(any(row["e_to_c_allowed"] for row in precursor_transport))

    def test_multistep_periplasm_transport_directions(self) -> None:
        model_path = self.directory / "periplasm.xml"
        model_path.write_text(SYNTHETIC_PERIPLASM_SBML, encoding="utf-8")
        model = load_model(model_path)
        transport = transport_records(model, "target")
        outward = [row for row in transport if row["c_to_e_allowed"]]
        inward = [row for row in transport if row["e_to_c_allowed"]]
        self.assertEqual(len(outward), 1)
        self.assertEqual(len(inward), 1)
        self.assertEqual(outward[0]["transport_path_compartments"], "c → p → e")
        self.assertIn("T_cp", outward[0]["transport_path_reactions"])
        self.assertIn("T_pe", outward[0]["transport_path_reactions"])

    def test_user_target_selection_is_explicit_and_traceable(self) -> None:
        target_file = self.directory / "targets.tsv"
        target_file.write_text(
            TARGET_HEADER
            + "Example target\ttarget\tPain biology\tNeuroactive metabolite\t"
            + "Reported microbial metabolite\tA priori literature selection\tdoi:example\n",
            encoding="utf-8",
        )
        result = validate_targets(self.directory, target_file)
        self.assertEqual(result["summary"]["user_selected_targets"], 1)
        self.assertEqual(result["summary"]["targets_with_gem_identifier"], 1)
        self.assertEqual(result["rows"][0]["decision"], "Keep")
        self.assertIn("does not screen all metabolites", result["metadata"]["selection_rule"])

    def test_unresolved_user_target_is_reported(self) -> None:
        target_file = self.directory / "targets.tsv"
        target_file.write_text(
            TARGET_HEADER
            + "Missing target\tmissing\tPain biology\tOther\tMicrobial evidence\t"
            + "A priori literature selection\tdoi:example\n",
            encoding="utf-8",
        )
        result = validate_targets(self.directory, target_file)
        self.assertEqual(result["rows"][0]["decision"], "Unresolved")
        self.assertEqual(result["summary"]["unresolved_targets"], 1)

    def test_duplicate_target_identifier_is_rejected(self) -> None:
        target_file = self.directory / "targets.tsv"
        row = (
            "First\ttarget\tPain\tOther\tMicrobial evidence\tLiterature\tdoi:1\n"
            "Second\ttarget\tPain\tOther\tMicrobial evidence\tLiterature\tdoi:2\n"
        )
        target_file.write_text(TARGET_HEADER + row, encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "Duplicate base_id"):
            read_targets(target_file)

    def test_expected_answer_column_is_rejected_from_discovery(self) -> None:
        target_file = self.directory / "targets.tsv"
        target_file.write_text(
            TARGET_HEADER.rstrip("\n")
            + "\texpected_precursor_id\n"
            + "Example target\ttarget\tPain\tOther\tMicrobial evidence\t"
            + "Literature\tdoi:example\tprec\n",
            encoding="utf-8",
        )
        with self.assertRaisesRegex(ValueError, "validation-only expected-answer"):
            read_targets(target_file)

    def test_executable_path_uses_platform_separator(self) -> None:
        with mock.patch("step4_intracellular_production.os.pathsep", ";"):
            with mock.patch.dict(os.environ, {"PATH": r"C:\\Windows\\System32"}):
                environment = executable_environment(Path("/portable/mene"))
        self.assertEqual(
            environment["PATH"],
            r"/portable;C:\\Windows\\System32",
        )

    def test_menetools_unproducible_target_is_a_null_result(self) -> None:
        executable = self.directory / "mene"
        executable.write_text("", encoding="utf-8")
        network = self.directory / "network.xml"
        seeds = self.directory / "seeds.xml"
        target = self.directory / "target.xml"
        output = self.directory / "output.json"
        network.write_text("<sbml/>", encoding="utf-8")
        seeds.write_text("<sbml/>", encoding="utf-8")
        target.write_text(
            '<sbml xmlns="http://www.sbml.org/sbml/level3/version1/core">'
            '<model id="target"><listOfSpecies>'
            '<species id="M_missing__91__c__93__" compartment="c"/>'
            '</listOfSpecies></model></sbml>',
            encoding="utf-8",
        )
        completed = mock.Mock(
            returncode=1,
            stdout=(
                "1 unproducible targets:\nM_missing__91__c__93__\n"
                "UnboundLocalError: one_path"
            ),
        )
        with mock.patch(
            "step4_intracellular_production.subprocess.run",
            return_value=completed,
        ):
            result = run_menetools(executable, network, seeds, target, output)
        self.assertEqual(
            result["unproducible_targets_lst"],
            ["M_missing__91__c__93__"],
        )
        self.assertTrue(result["menetools_compatibility_fallback"])

    def test_net_positive_target_path_is_retained(self) -> None:
        model = load_model(self.model_path)
        records = [
            record
            for record in directional_reaction_records(model)
            if record["reaction_id"] in {"T_prec", "PROD"}
        ]
        result = net_production_feasibility(
            records,
            {"M_prec__91__e__93__"},
            "M_target__91__c__93__",
        )
        self.assertTrue(result["supported"])
        self.assertGreater(result["net_target"], 0)

    def test_generic_redox_supported_target_path_is_retained(self) -> None:
        model_path = self.directory / "redox.xml"
        model_path.write_text(SYNTHETIC_REDOX_SBML, encoding="utf-8")
        model = load_model(model_path)
        records = directional_reaction_records(model)
        result = net_production_feasibility(
            records,
            {
                "M_prec__91__e__93__",
                "M_nad__91__c__93__",
                "M_so3__91__c__93__",
            },
            "M_target__91__c__93__",
        )
        self.assertTrue(result["supported"])
        self.assertGreater(result["net_target"], 0)

    def test_zero_net_target_cycle_is_rejected(self) -> None:
        model_path = self.directory / "target_cycle.xml"
        model_path.write_text(SYNTHETIC_TARGET_CYCLE_SBML, encoding="utf-8")
        model = load_model(model_path)
        records = [
            record
            for record in directional_reaction_records(model)
            if record["reaction_id"]
            in {"T_prec", "TARGET_CONSUMING_STEP", "TARGET_REFORMING_STEP"}
        ]
        result = net_production_feasibility(
            records,
            {"M_prec__91__e__93__"},
            "M_target__91__c__93__",
        )
        self.assertFalse(result["supported"])
        self.assertIn(
            "TARGET_CONSUMING_STEP__forward",
            result["target_consuming_reactions"],
        )


class ReleaseMetadataTest(unittest.TestCase):
    def test_menetools_is_found_beside_symlinked_virtualenv_python(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            bin_directory = Path(temporary) / "bin"
            bin_directory.mkdir()
            python_link = bin_directory / "python"
            python_link.symlink_to(Path(sys.executable).resolve())
            mene = bin_directory / "mene"
            mene.write_text("", encoding="utf-8")
            with mock.patch.object(find_crossfeeding.sys, "executable", str(python_link)):
                with mock.patch("find_crossfeeding.shutil.which", return_value=None):
                    self.assertEqual(find_crossfeeding.resolve_mene(), mene.resolve())

    def test_scientific_stages_use_fixed_python_hash_seed(self) -> None:
        with mock.patch("find_crossfeeding.subprocess.run") as run:
            run.return_value.returncode = 0
            find_crossfeeding.run_stage("stage.py", [])
        self.assertEqual(run.call_args.kwargs["env"]["PYTHONHASHSEED"], "0")

    def test_public_python_workbook_writer(self) -> None:
        from openpyxl import load_workbook

        result = {
            "metadata": {"analysis": "test"},
            "summary": {"kept": 1},
            "rows": [{"target_name": "GABA", "decision": "Keep"}],
            "evidence": [{"reaction_id": "R_TEST"}],
        }
        with tempfile.TemporaryDirectory() as temporary:
            output = Path(temporary) / "result.xlsx"
            write_workbook(result, output, "step2")
            workbook = load_workbook(output, read_only=True)
            self.assertEqual(
                workbook.sheetnames,
                ["Run Summary", "Results", "Evidence"],
            )

    def test_python_dependency_files_are_aligned(self) -> None:
        requirements = {
            line.strip()
            for line in (RELEASE_DIRECTORY / "requirements.txt")
            .read_text(encoding="utf-8")
            .splitlines()
            if line.strip() and not line.lstrip().startswith("#")
        }
        environment = (RELEASE_DIRECTORY / "environment.yml").read_text(
            encoding="utf-8"
        )
        for requirement in requirements:
            self.assertIn(f"- {requirement}", environment)

    def test_model_manifest_has_explicit_species_and_checksums(self) -> None:
        manifest = read_model_manifest(RELEASE_DIRECTORY / "config" / "model_manifest.tsv")
        self.assertEqual(len(manifest), 5)
        for record in manifest.values():
            self.assertTrue(record["species"])
            self.assertEqual(len(record["sha256"]), 64)

    def test_manuscript_edge_regression_file(self) -> None:
        expected_path = RELEASE_DIRECTORY / "config" / "expected_manuscript_edges.tsv"
        with expected_path.open(encoding="utf-8", newline="") as handle:
            expected = list(csv.DictReader(handle, delimiter="\t"))
        self.assertEqual(len(expected), 10)
        self.assertEqual(len({tuple(row.values()) for row in expected}), 10)

    def test_saved_validation_results_match_expected_cases(self) -> None:
        expected_paths = sorted(
            (RELEASE_DIRECTORY / "validation" / "cases").glob(
                "*/config/expected_result.tsv"
            )
        )
        cases = []
        for expected_path in expected_paths:
            with expected_path.open(encoding="utf-8", newline="") as handle:
                cases.extend(csv.DictReader(handle, delimiter="\t"))
        self.assertEqual(len(cases), 3)
        self.assertSetEqual(
            {row["case_id"] for row in cases},
            {
                "positive_succinate_propionate",
                "positive_ncc2705_rectale",
                "literature_positive_lactate_h2s",
            },
        )
        self.assertTrue(all("failed_tests" not in row["result_json"] for row in cases))
        for case in cases:
            result_path = RELEASE_DIRECTORY / case["result_json"]
            self.assertTrue(result_path.is_file(), case["case_id"])
            result = json.loads(result_path.read_text(encoding="utf-8"))
            expected_edge = (
                case["donor_model_id"],
                case["precursor_id"],
                case["recipient_model_id"],
                case["final_target_id"],
            )
            retained = {
                (
                    str(row["donor_model_id"]),
                    str(row["precursor_id"]),
                    str(row["recipient_model_id"]),
                    str(row["final_target_id"]),
                )
                for row in result["rows"]
                if row["decision"] == "Keep"
            }
            self.assertIn(expected_edge, retained, case["case_id"])
            expected_audits = [
                row
                for row in result.get("donor_audits", [])
                if row.get("donor_model_id") == case["donor_model_id"]
                and row.get("precursor_id") == case["precursor_id"]
            ]
            self.assertEqual(len(expected_audits), 1, case["case_id"])
            self.assertTrue(expected_audits[0]["donor_pass"], case["case_id"])
            self.assertFalse(
                expected_audits[0].get("menetools_compatibility_fallback"),
                case["case_id"],
            )

    def test_expected_answers_are_compared_only_after_prediction(self) -> None:
        comparisons = validate_expected_cases.compare_expected_cases()
        self.assertEqual(len(comparisons), 7)
        self.assertTrue(all(row["pass"] == "PASS" for row in comparisons))

    def test_reduced_redox_carriers_are_not_bootstrapped(self) -> None:
        self.assertTrue({"nad", "nadp", "fad"} <= COFACTOR_BOOTSTRAP)
        self.assertTrue({"nadh", "nadph", "fadh2"}.isdisjoint(COFACTOR_BOOTSTRAP))

    def test_matched_case1_has_one_positive_and_one_negative_donor(self) -> None:
        case = "case1_two_donors_lactate_butyrate"
        result_path = (
            RELEASE_DIRECTORY
            / "validation"
            / "cases"
            / "01_two_donors_control"
            / "results"
            / f"{case}.json"
        )
        result = json.loads(result_path.read_text(encoding="utf-8"))
        rows = {
            (row["donor_model_id"], row["recipient_model_id"]): row
            for row in result["rows"]
            if row["precursor_id"] == "lac_L" and row["final_target_id"] == "but"
        }
        positive = rows[
            (
                "M_Bifidobacterium_adolescentis_L2_32",
                "M_Eubacterium_hallii_L2_7",
            )
        ]
        negative = rows[
            (
                "M_Paraprevotella_xylaniphila_YIT_11841",
                "M_Eubacterium_hallii_L2_7",
            )
        ]
        self.assertEqual(positive["decision"], "Keep")
        self.assertEqual(negative["decision"], "Exclude")
        self.assertFalse(negative["internal_production"])
        self.assertGreaterEqual(result["summary"]["potential_edges_kept"], 1)

    def test_matched_case2_has_one_positive_and_one_negative_consumer(self) -> None:
        case = "case2_two_consumers_lactate_butyrate"
        result_path = (
            RELEASE_DIRECTORY
            / "validation"
            / "cases"
            / "02_two_consumers_control"
            / "results"
            / f"{case}.json"
        )
        result = json.loads(result_path.read_text(encoding="utf-8"))
        retained = {
            (
                row["donor_model_id"],
                row["precursor_id"],
                row["recipient_model_id"],
                row["final_target_id"],
            )
            for row in result["rows"]
            if row["decision"] == "Keep"
        }
        positive = (
            "M_Bifidobacterium_adolescentis_L2_32",
            "lac_L",
            "M_Eubacterium_hallii_L2_7",
            "but",
        )
        negative = (
            "M_Bifidobacterium_adolescentis_L2_32",
            "lac_L",
            "M_Phascolarctobacterium_succinatutens_YIT_12067",
            "but",
        )
        self.assertIn(positive, retained)
        self.assertNotIn(negative, retained)

    @unittest.skipUnless(
        os.environ.get("CROSSFEEDING_RESULT_JSON"),
        "set CROSSFEEDING_RESULT_JSON for result regression",
    )
    def test_external_result_matches_expected_edges(self) -> None:
        result_path = Path(os.environ["CROSSFEEDING_RESULT_JSON"]).expanduser().resolve()
        result = json.loads(result_path.read_text(encoding="utf-8"))
        observed = {
            (
                str(row["donor_model_id"]),
                str(row["precursor_id"]),
                str(row["recipient_model_id"]),
                str(row["final_target_id"]),
            )
            for row in result["rows"]
            if row["decision"] == "Keep"
        }
        expected_path = RELEASE_DIRECTORY / "config" / "expected_manuscript_edges.tsv"
        with expected_path.open(encoding="utf-8", newline="") as handle:
            expected = {
                (
                    row["donor_model_id"],
                    row["precursor_id"],
                    row["recipient_model_id"],
                    row["final_target_id"],
                )
                for row in csv.DictReader(handle, delimiter="\t")
            }
        self.assertSetEqual(observed, expected)


if __name__ == "__main__":
    unittest.main()
