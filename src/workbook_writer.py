"""Write publication-facing Excel workbooks using the public openpyxl package."""

from __future__ import annotations

from collections.abc import Iterable, Mapping
from pathlib import Path
from typing import Any


INVALID_SHEET_CHARACTERS = set("[]:*?/\\")


def _cell_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, Mapping):
        return "; ".join(f"{key}={_cell_value(item)}" for key, item in value.items())
    if isinstance(value, Iterable):
        return "; ".join(str(_cell_value(item)) for item in value)
    return str(value)


def _sheet_name(value: str) -> str:
    cleaned = "".join("_" if char in INVALID_SHEET_CHARACTERS else char for char in value)
    return cleaned[:31] or "Sheet"


def _write_key_values(sheet, title: str, values: Mapping[str, Any], start_row: int) -> int:
    from openpyxl.styles import Font

    sheet.cell(start_row, 1, title)
    sheet.cell(start_row, 1).font = Font(bold=True, size=12)
    row = start_row + 1
    for key, value in values.items():
        sheet.cell(row, 1, str(key))
        sheet.cell(row, 2, _cell_value(value))
        row += 1
    return row


def _write_table(workbook, name: str, records: list[dict[str, Any]]) -> None:
    from openpyxl.styles import Font, PatternFill

    sheet = workbook.create_sheet(_sheet_name(name))
    if not records:
        sheet.append(["No records"])
        return
    headers: list[str] = []
    for record in records:
        for key in record:
            if key not in headers:
                headers.append(key)
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    for record in records:
        sheet.append([_cell_value(record.get(header)) for header in headers])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        letter = column[0].column_letter
        width = max(len(str(cell.value or "")) for cell in column)
        sheet.column_dimensions[letter].width = min(max(width + 2, 10), 55)


def write_workbook(result: dict[str, Any], output_path: Path, kind: str) -> None:
    try:
        from openpyxl import Workbook
    except ImportError as error:
        raise RuntimeError(
            "Excel output requires openpyxl. Install the release environment or "
            "run with JSON output only."
        ) from error

    output_path = output_path.expanduser().resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Run Summary"
    next_row = _write_key_values(summary_sheet, "Metadata", result.get("metadata", {}), 1)
    _write_key_values(summary_sheet, "Summary", result.get("summary", {}), next_row + 1)
    summary_sheet.column_dimensions["A"].width = 38
    summary_sheet.column_dimensions["B"].width = 95
    summary_sheet.freeze_panes = "A2"

    if result.get("strain_key"):
        _write_table(workbook, "Strain Key", list(result["strain_key"]))
    _write_table(workbook, "Results", list(result.get("rows", [])))
    if result.get("pathways"):
        _write_table(workbook, "Candidate Pathways", list(result["pathways"]))
    if result.get("donor_audits"):
        _write_table(workbook, "Donor Audits", list(result["donor_audits"]))
    if result.get("evidence"):
        evidence_name = "Donor Evidence" if kind == "step6" else "Evidence"
        _write_table(workbook, evidence_name, list(result["evidence"]))
    if result.get("consumer_evidence"):
        _write_table(workbook, "Consumer Evidence", list(result["consumer_evidence"]))
    if result.get("consumer_audits"):
        _write_table(workbook, "Consumer Audits", list(result["consumer_audits"]))
    workbook.save(output_path)
